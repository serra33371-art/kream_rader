"""네이버 크리에이터 어드바이저(블로그 통계) 데이터를 긁어서 CSV / 구글 시트로 저장.

사용법 요약 (자세한 건 README.md):
  1. 크롬 개발자도구 > Network > Fetch/XHR 에서 통계 API 요청을
     우클릭 > Copy > Copy as cURL (bash) 로 복사해 curl.txt 에 붙여넣기
  2. python scrape_stats.py --start 2026-09-01 --end 2026-09-22
  3. (선택) --sheet-id 를 주면 구글 시트에도 바로 붙여넣음
"""
import argparse
import base64
import csv
import json
import re
import shlex
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import requests

try:
    # 회사망처럼 보안 프로그램이 HTTPS 를 검사하는 환경에서는 윈도우 인증서 저장소를
    # 써야 SSL 오류(CERTIFICATE_VERIFY_FAILED)가 안 난다. (pip install truststore)
    import truststore
    truststore.inject_into_ssl()
except ImportError:
    pass

DATE_RE = re.compile(r'^\d{4}-\d{2}-\d{2}$')
# 이 헤더들은 requests 가 알아서 처리하므로 복사본에서 제외
SKIP_HEADERS = {'content-length', 'accept-encoding', 'host',
                # 캐시 헤더가 있으면 서버가 304(빈 응답)를 돌려준다
                'if-none-match', 'if-modified-since'}


def parse_curl(text):
    """'Copy as cURL (bash)' 문자열에서 url / headers / cookies 를 뽑아낸다."""
    text = text.replace('\\\n', ' ').replace('\\\r\n', ' ')
    tokens = shlex.split(text)
    url, headers, cookies = None, {}, {}
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if tok in ('-H', '--header'):
            key, _, value = tokens[i + 1].partition(':')
            key, value = key.strip(), value.strip()
            if key.lower() == 'cookie':
                cookies.update(_parse_cookie(value))
            elif key.lower() not in SKIP_HEADERS:
                headers[key] = value
            i += 2
        elif tok in ('-b', '--cookie'):
            cookies.update(_parse_cookie(tokens[i + 1]))
            i += 2
        elif tok in ('-X', '--request', '--data', '--data-raw', '--data-binary', '-d'):
            i += 2
        elif tok.startswith('http'):
            url = tok
            i += 1
        else:
            i += 1
    if not url:
        sys.exit('curl.txt 에서 URL 을 찾지 못했습니다. "Copy as cURL (bash)" 로 복사했는지 확인하세요.')
    return url, headers, cookies


def _parse_cookie(raw):
    out = {}
    for part in raw.split(';'):
        if '=' in part:
            k, v = part.strip().split('=', 1)
            out[k] = v
    return out


def with_date(url, day, start_key, end_key):
    """URL 의 날짜 파라미터를 day 로 바꾼다.

    --start-key/--end-key 를 주면 그 파라미터만, 아니면 YYYY-MM-DD 형태인
    모든 파라미터를 바꾼다.
    """
    parts = urlsplit(url)
    params = parse_qsl(parts.query, keep_blank_values=True)
    ds = day.isoformat()
    new = []
    for k, v in params:
        if start_key or end_key:
            if k in (start_key, end_key):
                v = ds
        elif DATE_RE.match(v):
            v = ds
        new.append((k, v))
    return urlunsplit(parts._replace(query=urlencode(new)))


def find_rows(obj):
    """응답 JSON 안에서 가장 큰 'dict 리스트'를 찾아 표 데이터로 사용."""
    best = []

    def walk(o):
        nonlocal best
        if isinstance(o, list):
            if o and all(isinstance(x, dict) for x in o) and len(o) > len(best):
                best = o
            for x in o:
                walk(x)
        elif isinstance(o, dict):
            for v in o.values():
                walk(v)

    walk(obj)
    if not best and isinstance(obj, dict):
        best = [obj]
    return best


def flatten(d, prefix=''):
    out = {}
    for k, v in d.items():
        key = f'{prefix}{k}'
        if isinstance(v, dict):
            out.update(flatten(v, key + '.'))
        elif isinstance(v, list):
            out[key] = json.dumps(v, ensure_ascii=False)
        else:
            out[key] = v
    return out


LINK_KEYS = ('url', 'link', 'contentUrl', 'postUrl')
ID_KEYS = ('logNo', 'contentId', 'postId', 'documentId')


def add_blog_link(row, channel_id):
    """응답에 글 주소가 없으면 blogId + 글번호(logNo 등)로 링크를 만들어 넣는다."""
    if not channel_id or any(row.get(k) for k in LINK_KEYS):
        return row
    for k in ID_KEYS:
        for col, v in row.items():
            if col.split('.')[-1] == k and str(v).isdigit():
                return {**row, 'link': f'https://blog.naver.com/{channel_id}/{v}'}
    return row


def daterange(start, end, step='day'):
    if step == 'month':
        d = start.replace(day=1)
        while d <= end:
            yield d
            d = (d + timedelta(days=32)).replace(day=1)
        return
    d = start
    while d <= end:
        yield d
        d += timedelta(days=1)


def write_xlsx(rows, columns, path):
    """'전체' 시트 + 기간(월)별 시트로 나눠 엑셀 파일 저장."""
    from openpyxl import Workbook  # pip install openpyxl

    wb = Workbook()
    groups = {'전체': rows}
    if '기간' in columns:
        for r in rows:
            groups.setdefault(str(r['기간']), []).append(r)
    for i, (name, group) in enumerate(groups.items()):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = re.sub(r'[\\/*?:\[\]]', '_', name)[:31]
        ws.append(columns)
        for r in group:
            ws.append([r.get(c, '') for c in columns])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col_idx, c in enumerate(columns, 1):
            width = max([len(str(c))] + [len(str(r.get(c, ''))) for r in group[:200]])
            ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = min(width + 2, 60)
            if c == 'link':
                for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    if row[0].value:
                        row[0].hyperlink = row[0].value
                        row[0].style = 'Hyperlink'
    wb.save(path)
    print(f'{path} 저장 완료 (시트 {len(groups)}개: 전체 + 기간별)')


def upload_to_sheet(rows, columns, sheet_id, worksheet, creds_path):
    import gspread  # pip install gspread

    gc = gspread.service_account(filename=creds_path)
    sh = gc.open_by_key(sheet_id)
    try:
        ws = sh.worksheet(worksheet)
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title=worksheet, rows=1000, cols=len(columns))
    values = [columns] + [[r.get(c, '') for c in columns] for r in rows]
    ws.clear()
    ws.update(values, 'A1')
    print(f'구글 시트 "{worksheet}" 탭에 {len(rows)}행 업로드 완료')


def iter_curl(args):
    url, headers, cookies = parse_curl(Path(args.curl).read_text(encoding='utf-8'))
    session = requests.Session()
    session.headers.update(headers)
    session.cookies.update(cookies)

    if args.start:
        start = date.fromisoformat(args.start)
        end = date.fromisoformat(args.end) if args.end else date.today() - timedelta(days=1)
        targets = [(d.isoformat()[:7] if args.step == 'month' else d.isoformat(), with_date(url, d, args.start_key, args.end_key))
                   for d in daterange(start, end, args.step)]
    else:
        targets = [('', url)]

    channel_id = dict(parse_qsl(urlsplit(url).query)).get('channelId')
    for n, (label, target) in enumerate(targets):
        if n:
            time.sleep(args.delay)
        resp = session.get(target, timeout=20)
        if resp.status_code in (401, 403):
            sys.exit(f'[{resp.status_code}] 인증 실패 - 쿠키가 만료됐거나 서명(x-ca-sig)이 막혔습니다.\n'
                     '개발자도구에서 cURL 을 다시 복사하거나, HAR 방식(--har)을 쓰세요.')
        resp.raise_for_status()
        try:
            data = resp.json()
        except ValueError:
            sys.exit('JSON 이 아닌 응답입니다. Fetch/XHR 요청(api 주소)을 복사했는지 확인하세요.\n'
                     + resp.text[:300])
        yield label, channel_id, data


def iter_har(path, url_filter):
    """개발자도구 Network > 'Export HAR' 로 저장한 파일에서 응답을 꺼낸다.

    화면에서 월/날짜를 넘겨가며 찍힌 요청을 그대로 쓰므로 서명·쿠키 문제가 없다.
    """
    har = json.loads(Path(path).read_text(encoding='utf-8'))
    seen = set()
    for entry in har['log']['entries']:
        req_url = entry['request']['url']
        if url_filter not in req_url:
            continue
        content = entry['response'].get('content', {})
        text = content.get('text')
        if not text:
            continue  # 304 등 본문 없는 응답 (Disable cache 체크 필요)
        if content.get('encoding') == 'base64':
            text = base64.b64decode(text).decode('utf-8')
        params = dict(parse_qsl(urlsplit(req_url).query))
        label = params.get('date') or params.get('startDate') or ''
        interval = params.get('interval')
        if interval == 'month' and label:
            label = label[:7]  # 2025-09-01 -> 2025-09
        elif interval:
            label = f'{label} ({interval})'  # 일간/주간 조회분은 월간과 구분
        if label in seen:
            continue
        seen.add(label)
        yield label, params.get('channelId'), json.loads(text)


def main():
    p = argparse.ArgumentParser(description='네이버 블로그 통계 스크래퍼')
    p.add_argument('--curl', default='curl.txt', help='Copy as cURL(bash) 내용을 저장한 파일')
    p.add_argument('--har', help='curl 대신 개발자도구에서 저장한 .har 파일 사용')
    p.add_argument('--har-filter', default='cv-ranks', help='HAR 에서 골라낼 요청 주소 일부')
    p.add_argument('--start', help='시작일 YYYY-MM-DD (생략하면 복사한 URL 그대로 1번만 호출)')
    p.add_argument('--end', help='종료일 YYYY-MM-DD (기본: 어제)')
    p.add_argument('--start-key', help='시작일 파라미터 이름 (예: startDate). 생략 시 자동 감지')
    p.add_argument('--end-key', help='종료일 파라미터 이름 (예: endDate)')
    p.add_argument('--step', choices=['day', 'month'], default='day',
                   help='day: 하루씩 / month: 매월 1일로 (조회수 순위 같은 월간 화면)')
    p.add_argument('--out', default='blog_stats.csv', help='저장할 CSV 경로')
    p.add_argument('--xlsx', help='엑셀 파일로도 저장 (전체 + 월별 시트). 예: cv_ranks.xlsx')
    p.add_argument('--dump', action='store_true', help='원본 JSON 을 raw/ 폴더에 저장 (구조 확인용)')
    p.add_argument('--delay', type=float, default=1.0, help='요청 간 대기(초)')
    p.add_argument('--sheet-id', help='구글 시트 ID (URL 의 /d/<ID>/edit 부분)')
    p.add_argument('--worksheet', default='blog_stats', help='구글 시트 탭 이름')
    p.add_argument('--creds', default='service_account.json', help='구글 서비스계정 키 파일')
    args = p.parse_args()

    if args.har:
        source = iter_har(args.har, args.har_filter)
    else:
        source = iter_curl(args)

    all_rows = []
    for label, channel_id, data in source:
        if args.dump:
            Path('raw').mkdir(exist_ok=True)
            name = label or datetime.now().strftime('%Y%m%d_%H%M%S')
            Path(f'raw/{name}.json').write_text(
                json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

        rows = [flatten(r) for r in find_rows(data)]
        for r in rows:
            r = add_blog_link(r, channel_id)
            if label:
                r = {'기간': label, **r}
            all_rows.append(r)
        print(f'{label or "요청"}: {len(rows)}행')

    if not all_rows:
        sys.exit('가져온 데이터가 없습니다. --dump 로 원본 JSON 을 확인하세요.')

    columns = []
    for r in all_rows:
        for k in r:
            if k not in columns:
                columns.append(k)

    # utf-8-sig: 엑셀/구글시트에서 한글 안 깨지게
    with open(args.out, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=columns)
        w.writeheader()
        w.writerows(all_rows)
    print(f'{args.out} 저장 완료 ({len(all_rows)}행, {len(columns)}열)')

    if args.xlsx:
        write_xlsx(all_rows, columns, args.xlsx)

    if args.sheet_id:
        upload_to_sheet(all_rows, columns, args.sheet_id, args.worksheet, args.creds)


if __name__ == '__main__':
    main()
